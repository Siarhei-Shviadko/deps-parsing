import io

import olefile
import openpyxl as xl
from deps_tabular_layout.models import Cell, SheetBuilder, TableBuilder, TabularLayout
from deps_tabular_layout.models.tabular_layout.sheet.abstract_builder import (
    AbstractSheetBuilder,
)
from xlrd import open_workbook

from ...excel import (
    CommentMapper,
    DataTypeMapper,
    FormatMapper,
    LineBasedDetectionStrategy,
    PositionMapper,
)
from ...table_reference import TableReference
from ..abstract_parser import AbstractTabularLayoutParser

__all__ = ["ExcelParser"]


class ExcelParser(AbstractTabularLayoutParser):
    _CELL_BATCH_SIZE = 200

    def _perform_parsing(self, blob: bytes, tabular_layout: TabularLayout) -> TabularLayout:
        if self._is_xls(blob):
            blob = self._convert_xls_to_xlsx(blob)

        with io.BytesIO(blob) as bytes_stream:
            workbook = xl.load_workbook(bytes_stream, data_only=True)

            builder = tabular_layout
            for sheet in workbook:
                builder = self._record_sheet(layout_id=tabular_layout.id(), sheet=sheet, tl_builder=builder)

        builder.build()
        return tabular_layout

    def _record_sheet(
        self,
        layout_id: str,
        sheet: xl.worksheet.worksheet.Worksheet,
        tl_builder: TabularLayout,
    ) -> SheetBuilder:
        tl_builder = (
            tl_builder.with_sheet().with_title(sheet.title).with_visibility(is_hidden=sheet.sheet_state == "hidden")
        )

        tables: list[TableReference] = self._detect_tables(sheet)

        for table_reference in tables:
            tl_builder = self._record_table(
                layout_id=layout_id,
                sheet=sheet,
                table_reference=table_reference,
                builder=tl_builder,
            )

        return tl_builder

    def _record_table(
        self,
        layout_id: str,
        sheet: xl.worksheet.worksheet.Worksheet,
        table_reference: TableReference,
        builder: SheetBuilder,
    ) -> AbstractSheetBuilder:
        builder = builder.with_table(
            left_top_corner=table_reference.top_left,
            right_bottom_corner=table_reference.bottom_right,
        )

        cells_batch: list[Cell] = []

        for row in sheet.iter_rows(
            # we are adding 1 because indexation in openpyxl starts from 1
            min_row=table_reference.min_row + 1,
            max_row=table_reference.max_row + 1,
            min_col=table_reference.min_column + 1,
            max_col=table_reference.max_column + 1,
        ):
            for cell in row:
                # we skip cells that are part of merged cells
                # and count only the top-left cell of the merged range
                if self._part_of_merged_cell(cell, sheet):
                    continue

                cells_batch.append(
                    self._record_cell(
                        cell=cell,
                        sheet=sheet,
                        builder=builder,
                    ),
                )

                if len(cells_batch) == self._CELL_BATCH_SIZE:
                    self._cell_repository.save_batch(layout_id, cells_batch)
                    cells_batch.clear()

        if cells_batch:
            self._cell_repository.save_batch(layout_id, cells_batch)

        return builder

    def _record_cell(
        self,
        cell: xl.cell.cell.Cell,
        sheet: xl.worksheet.worksheet.Worksheet,
        builder: TableBuilder,
    ) -> Cell:
        return builder.record_cell(
            content=str(cell.value) if cell.value is not None else "",
            position=PositionMapper.from_excel(
                cell=cell,
                sheet=sheet,
                table_builder=builder,
            ),
            data_type=DataTypeMapper.from_cell(cell),
            comment=CommentMapper.from_cell(cell),
            format_=FormatMapper.from_cell(cell),
        )

    def _detect_tables(self, sheet: xl.worksheet.worksheet.Worksheet) -> list[TableReference]:
        return LineBasedDetectionStrategy().detect_tables(sheet)

    def _part_of_merged_cell(self, cell: xl.cell.cell.Cell, sheet: xl.worksheet.worksheet.Worksheet) -> bool:
        if cell.coordinate not in sheet.merged_cells:
            return False

        return isinstance(cell, xl.cell.cell.MergedCell) and cell.value is None

    def _is_xls(self, blob: bytes) -> bool:
        try:
            with io.BytesIO(blob) as bytes_stream:
                if not olefile.isOleFile(bytes_stream):
                    return False
                ole = olefile.OleFileIO(bytes_stream)
                target_streams = [["Workbook"], ["Book"]]
                return any(s in ole.listdir() for s in target_streams)
        except Exception:
            return False

    def _convert_xls_to_xlsx(self, xls_blob: bytes) -> bytes:
        with io.BytesIO(xls_blob) as xls_io:
            book = open_workbook(file_contents=xls_io.read())

        xlsx_workbook = xl.Workbook()
        xlsx_workbook.remove(xlsx_workbook.active)

        for sheet_index in range(book.nsheets):
            xls_sheet = book.sheet_by_index(sheet_index)
            xlsx_sheet = xlsx_workbook.create_sheet(title=xls_sheet.name)

            for row_idx in range(xls_sheet.nrows):
                for col_idx in range(xls_sheet.ncols):
                    cell_value = xls_sheet.cell_value(row_idx, col_idx)
                    xlsx_sheet.cell(row=row_idx + 1, column=col_idx + 1, value=cell_value)

        with io.BytesIO() as output:
            xlsx_workbook.save(output)
            return output.getvalue()
